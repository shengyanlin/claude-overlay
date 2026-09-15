# -*- coding: utf-8 -*-
"""
The worker half of the 📎 attach-file feature: how a picked path becomes a content block.

This is the code that builds the actual API payload and parses real files on this machine,
and it shipped with no coverage at all while the same change deleted 469 lines of the
allowance gauge's tests. The UI-side tests (test_ui_attachments.py) stop at "the path
reached worker.ask" — everything below that line is here.

SAFETY: same contract as test_worker.py — construct ClaudeWorker(queue.Queue()) and call
pure/sync/static methods only. No network, no CLI, no GUI.
"""

import base64
import queue

import pytest

from docx import Document
from pptx import Presentation
from pptx.util import Emu

import worker as worker_module
from worker import ClaudeWorker
import config


def make_worker():
    return ClaudeWorker(queue.Queue())


def _drain(q):
    out = []
    while True:
        try:
            out.append(q.get_nowait())
        except queue.Empty:
            return out


@pytest.fixture
def docx_file(tmp_path):
    p = tmp_path / "report.docx"
    d = Document()
    d.add_paragraph("FIRST_PARAGRAPH")
    d.add_paragraph("")                      # empty paragraphs are skipped, not blank lines
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "CELL_LEFT"
    t.cell(0, 1).text = "CELL_RIGHT"
    d.save(p)
    return str(p)


@pytest.fixture
def pptx_file(tmp_path):
    """A deck with text at three depths: loose, inside a group, inside a nested group —
    plus a table and a speaker note."""
    p = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    loose = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(1000000), Emu(500000))
    loose.text_frame.text = "LOOSE_BOX"

    grp = slide.shapes.add_group_shape()
    inner = grp.shapes.add_textbox(Emu(0), Emu(600000), Emu(1000000), Emu(500000))
    inner.text_frame.text = "INSIDE_GROUP"
    nested = grp.shapes.add_group_shape()
    deep = nested.shapes.add_textbox(Emu(0), Emu(1200000), Emu(1000000), Emu(500000))
    deep.text_frame.text = "NESTED_TWICE"

    tbl = slide.shapes.add_table(1, 2, Emu(0), Emu(2000000), Emu(2000000), Emu(500000))
    tbl.table.cell(0, 0).text = "T_LEFT"
    tbl.table.cell(0, 1).text = "T_RIGHT"

    slide.notes_slide.notes_text_frame.text = "SPEAKER_NOTE"
    prs.save(p)
    return str(p)


@pytest.fixture
def png_file(tmp_path):
    from PIL import Image
    p = tmp_path / "shot.png"
    Image.new("RGB", (4, 4), (1, 2, 3)).save(p)
    return str(p)


@pytest.fixture
def pdf_file(tmp_path):
    p = tmp_path / "doc.pdf"
    p.write_bytes(b"%PDF-1.4\n%fake but non-empty\n")
    return str(p)


class TestExtractDocx:

    def test_paragraphs_and_table_cells(self, docx_file):
        text = ClaudeWorker._extract_docx(docx_file)
        assert "FIRST_PARAGRAPH" in text
        assert "CELL_LEFT | CELL_RIGHT" in text

    def test_empty_paragraphs_are_dropped(self, docx_file):
        assert "\n\n" not in ClaudeWorker._extract_docx(docx_file)

    def test_tables_stay_where_the_author_put_them(self, tmp_path):
        """doc.paragraphs then doc.tables is the obvious walk and it silently moves every
        table to the end, away from the prose that introduces it. "Prices" / table /
        "Prices exclude tax" would arrive with the qualifier attached to the wrong thing,
        and a model asked what the table shows answers from the wrong caption."""
        p = tmp_path / "ordered.docx"
        d = Document()
        d.add_paragraph("BEFORE_TABLE")
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).text = "IN_TABLE"
        d.add_paragraph("AFTER_TABLE")
        d.save(p)
        text = ClaudeWorker._extract_docx(str(p))
        assert text.index("BEFORE_TABLE") < text.index("IN_TABLE") < text.index("AFTER_TABLE")


class TestExtractPptx:

    def test_loose_text(self, pptx_file):
        assert "LOOSE_BOX" in ClaudeWorker._extract_pptx(pptx_file)

    def test_text_inside_a_group_survives(self, pptx_file):
        """A flat `for shape in slide.shapes` sees a GroupShape as one opaque item with no
        text frame, so everything inside it vanishes with no error — and a real deck groups
        constantly, which means the slides worth reading are the ones that come back
        emptiest. This is the regression that matters most in the whole file."""
        assert "INSIDE_GROUP" in ClaudeWorker._extract_pptx(pptx_file)

    def test_groups_nested_in_groups_survive_too(self, pptx_file):
        assert "NESTED_TWICE" in ClaudeWorker._extract_pptx(pptx_file)

    def test_tables_and_notes(self, pptx_file):
        text = ClaudeWorker._extract_pptx(pptx_file)
        assert "T_LEFT | T_RIGHT" in text
        assert "[Notes] SPEAKER_NOTE" in text

    def test_slides_are_numbered(self, pptx_file):
        assert "--- Slide 1 ---" in ClaudeWorker._extract_pptx(pptx_file)

    def test_a_deck_with_no_notes_is_not_given_any(self, tmp_path):
        """Reading .notes_slide CREATES a notes slide, so the has_notes_slide guard is load
        bearing: without it, extraction mutates the file's shape tree in memory and every
        slide grows an empty [Notes] line."""
        p = tmp_path / "bare.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(p)
        assert "[Notes]" not in ClaudeWorker._extract_pptx(str(p))


class TestDocTextBlock:

    def test_block_is_headed_with_the_filename(self, docx_file):
        w = make_worker()
        block, charged = w._doc_text_block(docx_file, ".docx")
        assert block["type"] == "text"
        assert block["text"].startswith("[Attached file: report.docx]")
        assert "FIRST_PARAGRAPH" in block["text"]

    def test_per_file_char_cap_truncates_and_says_so(self, docx_file, monkeypatch):
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_CHARS", 12)
        w = make_worker()
        block, charged = w._doc_text_block(docx_file, ".docx")
        assert "truncated" in block["text"]
        assert "this file is longer" in block["text"]

    def test_the_aggregate_budget_shrinks_the_room_left(self, docx_file, monkeypatch):
        """The per-file cap multiplies by the attachment count unless spend is carried
        across files. MAX_PENDING_IMAGES docs at the per-file limit is ~800k tokens."""
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_CHARS", 10_000)
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_TOTAL_CHARS", 20)
        w = make_worker()
        block, charged = w._doc_text_block(docx_file, ".docx", already=8)
        assert "truncated at 12 characters" in block["text"]

    def test_no_room_left_means_no_block(self, docx_file, monkeypatch):
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_TOTAL_CHARS", 100)
        w = make_worker()
        assert w._doc_text_block(docx_file, ".docx", already=100) == (None, 0)

    def test_an_empty_document_is_not_sent_as_a_blank_block(self, tmp_path):
        p = tmp_path / "blank.docx"
        Document().save(p)
        w = make_worker()
        assert w._doc_text_block(str(p), ".docx") == (None, 0)

    def test_a_corrupt_file_is_rejected_not_raised(self, tmp_path):
        p = tmp_path / "lying.docx"
        p.write_bytes(b"this is not a zip container")
        w = make_worker()
        assert w._doc_text_block(str(p), ".docx") == (None, 0)

    def test_the_size_cap_is_re_checked_here_not_only_in_the_picker(self, docx_file,
                                                                   monkeypatch):
        """The picker measured the file when it was CHOSEN; this runs later — after a queued
        message waits its turn, or after the file was replaced — and it is this call that
        unzips the container and parses every XML part. A .docx is a zip, so a few MB can
        expand to gigabytes of XML: the guard has to sit in front of the parser."""
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_BYTES", 1)
        w = make_worker()
        assert w._doc_text_block(docx_file, ".docx") == (None, 0)

    def test_a_missing_package_tells_the_user_how_to_fix_it(self, docx_file, monkeypatch):
        def no_docx(_p):
            raise ImportError("No module named 'docx'")

        monkeypatch.setattr(ClaudeWorker, "_extract_docx", staticmethod(no_docx))
        w = make_worker()
        assert w._doc_text_block(docx_file, ".docx") == (None, 0)
        kinds = _drain(w.ui)
        assert kinds and kinds[0][0] == "error"
        assert "update.cmd" in kinds[0][1]      # the command that actually installs it


class TestPdfBlock:

    def test_native_document_block(self, pdf_file):
        block, used = ClaudeWorker._pdf_block(pdf_file, 0)
        assert block["type"] == "document"
        assert block["source"]["media_type"] == "application/pdf"
        assert base64.b64decode(block["source"]["data"]).startswith(b"%PDF")
        assert used > 0

    def test_over_the_per_file_ceiling(self, pdf_file, monkeypatch):
        monkeypatch.setattr(worker_module, "MAX_INLINE_PDF_BYTES", 1)
        assert ClaudeWorker._pdf_block(pdf_file, 0) == (None, 0)

    def test_over_the_aggregate_ceiling(self, pdf_file, monkeypatch):
        monkeypatch.setattr(worker_module, "MAX_INLINE_TOTAL_BYTES", 10)
        assert ClaudeWorker._pdf_block(pdf_file, 9) == (None, 0)

    def test_an_empty_pdf_is_not_sent(self, tmp_path):
        p = tmp_path / "zero.pdf"
        p.write_bytes(b"")
        assert ClaudeWorker._pdf_block(str(p), 0) == (None, 0)

    def test_a_missing_file_is_rejected_not_raised(self, tmp_path):
        assert ClaudeWorker._pdf_block(str(tmp_path / "nope.pdf"), 0) == (None, 0)


class TestImageBlock:

    def test_media_type_follows_the_extension(self, tmp_path):
        from PIL import Image
        jpg = tmp_path / "x.jpg"
        Image.new("RGB", (2, 2)).save(jpg)
        block, used = ClaudeWorker._image_block(str(jpg), 0)
        assert block["source"]["media_type"] == "image/jpeg" and used > 0

    def test_an_unmapped_extension_defaults_to_png(self, tmp_path):
        """.bmp is an accepted image type with no entry in the media-type map, so it
        is the branch this covers. The first version passed a .png, which reaches the
        same default only because .png also happens to be unmapped - adding it to the
        map would have left the test green and the fallback untested."""
        from PIL import Image
        bmp = tmp_path / "x.bmp"
        Image.new("RGB", (2, 2)).save(bmp)
        block, _ = ClaudeWorker._image_block(str(bmp), 0)
        assert block["source"]["media_type"] == "image/png"


class TestBuildQueryRouting:
    """_build_query turns the flat path list the UI hands over into typed content blocks."""

    def _content(self, w, text, paths):
        coro = w._build_query(text, paths)
        assert not isinstance(coro, str), "inline mode should return a message generator"
        msgs = []

        async def collect():
            async for m in coro:
                msgs.append(m)

        import asyncio
        asyncio.run(collect())
        assert len(msgs) == 1
        return msgs[0]["message"]["content"]

    def test_each_kind_gets_its_own_block_type(self, w_inline, png_file, pdf_file, docx_file):
        content = self._content(w_inline, "look", [png_file, pdf_file, docx_file])
        types = [b["type"] for b in content]
        assert types.count("image") == 1
        assert types.count("document") == 1
        assert types.count("text") == 2      # the prompt itself + the extracted docx

    def test_an_unsupported_extension_is_counted_not_crashed(self, w_inline, tmp_path):
        p = tmp_path / "data.xlsx"
        p.write_bytes(b"not really a workbook")
        content = self._content(w_inline, "look", [str(p)])
        assert [b["type"] for b in content] == ["text"]       # prompt only
        errs = [e for e in _drain(w_inline.ui) if e[0] == "error"]
        assert errs and "isn't a supported type" in errs[0][1]

    def test_repeated_paths_are_deduped(self, w_inline, docx_file):
        content = self._content(w_inline, "look", [docx_file, docx_file])
        assert len([b for b in content if b["type"] == "text"]) == 2   # prompt + ONE docx

    def _docx(self, tmp_path, name, chars):
        p = tmp_path / name
        d = Document()
        d.add_paragraph("X" * chars)
        d.save(p)
        return str(p)

    def test_the_aggregate_doc_budget_is_carried_between_files(self, w_inline, tmp_path,
                                                               monkeypatch):
        """What the first document spends, the second does not get. Without the carry the
        per-file cap is multiplied by the attachment count — MAX_PENDING_IMAGES docs at the
        per-file limit is ~800k tokens, which no context survives."""
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_CHARS", 10_000)
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_TOTAL_CHARS", 400)
        small = self._docx(tmp_path, "small.docx", 100)
        big = self._docx(tmp_path, "big.docx", 5_000)
        content = self._content(w_inline, "look", [small, big])
        docs = [b["text"] for b in content if b["text"].startswith("[Attached file:")]
        assert len(docs) == 2
        assert "truncated" not in docs[0]          # small.docx fit whole
        assert "truncated" in docs[1]              # big.docx got only what was left
        # The budget counts PROSE, not the assembled block: the filename header and the
        # truncation note are fixed per-attachment overhead, bounded by the count cap.
        prose = 0
        for t in docs:
            body = t.split("]\n", 1)[1]
            prose += len(body.split("\n[... truncated")[0])
        assert prose <= 400        # the aggregate held, not 2x the per-file cap

    def test_a_document_that_no_budget_is_left_for_is_dropped_and_reported(
            self, w_inline, tmp_path, monkeypatch):
        """Budget fully spent by an earlier file: the later one is dropped, not sent as a
        zero-character block. Silently attaching an empty document would be worse than
        refusing it — the model would answer about a file it never actually received."""
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_CHARS", 10_000)
        monkeypatch.setattr(worker_module, "MAX_INLINE_DOC_TOTAL_CHARS", 200)
        first = self._docx(tmp_path, "first.docx", 5_000)
        second = self._docx(tmp_path, "second.docx", 5_000)
        content = self._content(w_inline, "look", [first, second])
        docs = [b["text"] for b in content if b["text"].startswith("[Attached file:")]
        assert len(docs) == 1 and "first.docx" in docs[0]
        errs = [e for e in _drain(w_inline.ui) if e[0] == "error"]
        assert errs and "Not sent" in errs[0][1] and "second.docx" in errs[0][1]

    def test_a_failed_file_does_not_spend_an_attachment_slot(self, w_inline, tmp_path,
                                                             monkeypatch):
        """The cap counts blocks BUILT, not paths seen. Counting paths let a file that
        produced nothing still spend a slot, so one corrupt document could evict a valid
        image later in the same turn and the message went out under its own cap."""
        monkeypatch.setattr(worker_module, "MAX_INLINE_IMAGES", 2)
        corrupt = tmp_path / "lying.docx"
        corrupt.write_bytes(b"not a zip")
        from PIL import Image
        imgs = []
        for n in (1, 2):
            q = tmp_path / f"i{n}.png"
            Image.new("RGB", (2, 2), (n, n, n)).save(q)
            imgs.append(str(q))
        content = self._content(w_inline, "look", [str(corrupt)] + imgs)
        assert len([b for b in content if b["type"] == "image"]) == 2

    def test_a_rejected_file_is_named_in_the_error(self, w_inline, tmp_path):
        """Same contract as the picker's rejections: the file and the cause, not a tally."""
        corrupt = tmp_path / "broken.docx"
        corrupt.write_bytes(b"not a zip")
        self._content(w_inline, "look", [str(corrupt)])
        errs = [e for e in _drain(w_inline.ui) if e[0] == "error"]
        assert errs and "broken.docx" in errs[0][1]
        assert "parsed" in errs[0][1]

    def test_over_the_count_cap_says_so_by_name(self, w_inline, tmp_path, monkeypatch):
        monkeypatch.setattr(worker_module, "MAX_INLINE_IMAGES", 1)
        from PIL import Image
        paths = []
        for n in (1, 2):
            q = tmp_path / f"i{n}.png"
            Image.new("RGB", (2, 2), (n, n, n)).save(q)
            paths.append(str(q))
        self._content(w_inline, "look", paths)
        errs = [e for e in _drain(w_inline.ui) if e[0] == "error"]
        assert errs and "i2.png" in errs[0][1] and "1-attachment limit" in errs[0][1]

    def test_the_count_cap_covers_documents_too(self, w_inline, tmp_path, monkeypatch):
        monkeypatch.setattr(worker_module, "MAX_INLINE_IMAGES", 2)
        paths = []
        for n in range(4):
            p = tmp_path / f"d{n}.docx"
            d = Document()
            d.add_paragraph(f"BODY_{n}")
            d.save(p)
            paths.append(str(p))
        content = self._content(w_inline, "look", paths)
        assert len([b for b in content if b["text"].startswith("[Attached file:")]) == 2


@pytest.fixture
def w_inline(monkeypatch):
    """A worker in inline-attachment mode, which is the default but is read from the module
    global — pin it so the legacy "Read the PNG path" mode can't silently take over."""
    monkeypatch.setattr(worker_module, "IMAGE_INPUT", "inline")
    return make_worker()
